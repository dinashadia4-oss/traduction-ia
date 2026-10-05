import os
import sqlite3
from docx import Document
from difflib import SequenceMatcher
from concurrent.futures import ThreadPoolExecutor, as_completed

DOSSIER = r"C:\Users\hatta\Desktop\serv"
BASE = "index_modeles_word.db"

MAX_WORKERS = 8


def trouver_docx():
    fichiers_word = []

    for racine, dossiers, fichiers in os.walk(DOSSIER):
        for fichier in fichiers:
            nom_lower = fichier.lower()

            if fichier.startswith("~$"):
                continue

            if nom_lower.endswith(".docx"):
                chemin = os.path.join(racine, fichier)
                fichiers_word.append(chemin)

    return fichiers_word


def lire_docx(chemin):
    try:
        doc = Document(chemin)
        texte = []

        for p in doc.paragraphs:
            if p.text.strip():
                texte.append(p.text.strip())

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        texte.append(cell.text.strip())

        contenu = "\n".join(texte).strip()
        return chemin, os.path.basename(chemin), contenu, None

    except Exception as e:
        return chemin, os.path.basename(chemin), "", str(e)


def creer_index():
    conn = sqlite3.connect(BASE)
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS fichiers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            chemin TEXT UNIQUE,
            nom TEXT,
            texte TEXT
        )
    """)

    print("\nRecherche dans tous les sous-dossiers...")
    fichiers = trouver_docx()

    print(f"Fichiers .docx trouvés : {len(fichiers)}")

    cur.execute("SELECT chemin FROM fichiers")
    deja_indexes = set(row[0] for row in cur.fetchall())

    fichiers_a_indexer = [f for f in fichiers if f not in deja_indexes]

    print(f"Déjà indexés : {len(deja_indexes)}")
    print(f"Nouveaux à indexer : {len(fichiers_a_indexer)}\n")

    total_ajoutes = 0
    total_erreurs = 0

    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        futures = [executor.submit(lire_docx, chemin) for chemin in fichiers_a_indexer]

        for i, future in enumerate(as_completed(futures), 1):
            chemin, nom, texte, erreur = future.result()

            if erreur:
                total_erreurs += 1
                print(f"[ERREUR] {nom}")
                continue

            cur.execute(
                "INSERT OR IGNORE INTO fichiers (chemin, nom, texte) VALUES (?, ?, ?)",
                (chemin, nom, texte[:50000])
            )

            total_ajoutes += 1

            if i % 50 == 0:
                conn.commit()
                print(f"Progression : {i}/{len(fichiers_a_indexer)}")

    conn.commit()
    conn.close()

    print("\nIndex terminé.")
    print(f"Nouveaux fichiers ajoutés : {total_ajoutes}")
    print(f"Fichiers avec erreur : {total_erreurs}")


def calculer_score(recherche, nom, texte):
    recherche = recherche.lower()
    nom = nom.lower()
    texte = texte.lower()

    score_final = 0

    if recherche in nom:
        score_final = max(score_final, 100)

    if recherche in texte:
        score_final = max(score_final, 95)

    score_nom = int(SequenceMatcher(None, recherche, nom).ratio() * 100)
    score_final = max(score_final, score_nom)

    mots = recherche.split()
    if mots:
        mots_trouves = sum(1 for mot in mots if mot in texte or mot in nom)
        score_mots = int((mots_trouves / len(mots)) * 90)
        score_final = max(score_final, score_mots)

    return score_final


def chercher():
    conn = sqlite3.connect(BASE)
    cur = conn.cursor()

    cur.execute("SELECT COUNT(*) FROM fichiers")
    total = cur.fetchone()[0]

    print(f"\nBase prête : {total} fichiers indexés.")
    print("Tape q pour quitter.\n")

    while True:
        recherche = input("Recherche modèle Word : ").strip()

        if recherche.lower() in ["q", "quit", "exit"]:
            break

        if not recherche:
            continue

        cur.execute("SELECT nom, chemin, texte FROM fichiers")
        resultats = []

        for nom, chemin, texte in cur.fetchall():
            score = calculer_score(recherche, nom, texte)

            if score >= 25:
                resultats.append((score, nom, chemin))

        resultats.sort(reverse=True)

        print("\nMeilleurs résultats :")

        if not resultats:
            print("Aucun résultat trouvé.")
        else:
            for i, (score, nom, chemin) in enumerate(resultats[:30], 1):
                print(f"\n{i}. {nom}")
                print(f"   Score : {score}%")
                print(f"   Chemin : {chemin}")

        print("\n" + "-" * 60)

    conn.close()


def afficher_statistiques():
    conn = sqlite3.connect(BASE)
    cur = conn.cursor()

    cur.execute("SELECT COUNT(*) FROM fichiers")
    total_index = cur.fetchone()[0]

    fichiers_reels = trouver_docx()

    print("\nStatistiques :")
    print(f"Fichiers .docx trouvés dans le dossier et sous-dossiers : {len(fichiers_reels)}")
    print(f"Fichiers indexés dans la base : {total_index}")

    conn.close()


def reindexer_tout():
    if os.path.exists(BASE):
        os.remove(BASE)
        print("Ancien index supprimé.")

    creer_index()


def menu():
    while True:
        print("\n==============================")
        print("RECHERCHE INTELLIGENTE WORD")
        print("==============================")
        print("1. Créer / mettre à jour l’index")
        print("2. Chercher un modèle Word")
        print("3. Voir statistiques")
        print("4. Réindexer tout depuis zéro")
        print("5. Quitter")

        choix = input("Choix : ").strip()

        if choix == "1":
            creer_index()
        elif choix == "2":
            chercher()
        elif choix == "3":
            afficher_statistiques()
        elif choix == "4":
            reindexer_tout()
        elif choix == "5":
            print("Terminé.")
            break
        else:
            print("Choix incorrect.")


if __name__ == "__main__":
    menu()